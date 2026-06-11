"""
utils/export.py
Exports resume to PDF (template-aware), DOCX, TXT.
"""
from __future__ import annotations
import io
import re


def is_bullet(line: str) -> bool:
    return line.startswith(("•", "-", "*", "▪", "▸", "◆"))


def clean_bullet_text(line: str) -> str:
    return re.sub(r"^[•\-\*▪▸◆\s]+", "", line).strip()


def strip_markdown(text: str) -> str:
    return re.sub(r"\*\*|__|\*|_", "", text)


def md_to_reportlab(text: str) -> str:
    # Replace **text** with <b>text</b>
    text = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", text)
    # Replace *text* with <i>text</i>
    text = re.sub(r"\*(.*?)\*", r"<i>\1</i>", text)
    # Replace __text__ with <b>text</b>
    text = re.sub(r"__(.*?)__", r"<b>\1</b>", text)
    # Replace _text_ with <i>text</i>
    text = re.sub(r"_(.*?)_", r"<i>\1</i>", text)
    return text


def clean_and_check_header(line: str) -> tuple[bool, str]:
    # Strip markdown symbols like *, #, _, -
    cleaned = re.sub(r"^[#\s\-*_*]+|[#\s\-*_*]+$", "", line).strip()
    if not cleaned:
        return False, ""
        
    common_headers = [
        "summary", "professional summary", "career summary", "summary of qualifications",
        "experience", "work experience", "professional experience", "employment history", "work history",
        "skills", "technical skills", "core competencies", "areas of expertise", "skills & tools",
        "education", "academic background", "certifications", "licenses & certifications",
        "projects", "personal projects", "key projects"
    ]
    
    if cleaned.lower() in common_headers:
        return True, cleaned.upper()
        
    if cleaned.isupper() and 3 <= len(cleaned) <= 50 and re.match(r"^[A-Z][A-Z\s\/&]{2,}$", cleaned):
        return True, cleaned
        
    return False, ""


def resume_to_txt_bytes(text: str) -> bytes:
    return text.encode("utf-8")


def text_to_docx_bytes(text: str) -> bytes:
    try:
        from docx import Document
        from docx.shared import Pt, RGBColor, Inches
        from docx.enum.text import WD_ALIGN_PARAGRAPH

        doc = Document()
        for section in doc.sections:
            section.top_margin    = Inches(0.75)
            section.bottom_margin = Inches(0.75)
            section.left_margin   = Inches(0.9)
            section.right_margin  = Inches(0.9)

        style = doc.styles["Normal"]
        style.font.name = "Calibri"
        style.font.size = Pt(10.5)

        name_done = False
        for line in text.split("\n"):
            stripped = line.strip()
            if not stripped:
                doc.add_paragraph()
                continue

            # Check for section header
            is_hdr, hdr_title = clean_and_check_header(stripped)
            if is_hdr:
                p   = doc.add_paragraph()
                run = p.add_run(hdr_title)
                run.bold = True
                run.font.size = Pt(11)
                run.font.color.rgb = RGBColor(0x1a, 0x50, 0xc8)
                p.paragraph_format.space_before = Pt(10)
                p.paragraph_format.space_after  = Pt(3)
                continue

            # Name
            if not name_done:
                name_candidate = strip_markdown(stripped).strip()
                if len(name_candidate.split()) <= 6 and not any(c in name_candidate for c in ["@", "|", "+"]):
                    p   = doc.add_paragraph()
                    run = p.add_run(name_candidate)
                    run.font.size = Pt(20)
                    run.bold = True
                    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    name_done = True
                    continue

            # Bullet
            if is_bullet(stripped):
                p = doc.add_paragraph(style="List Bullet")
                clean = clean_bullet_text(stripped)
                p.add_run(strip_markdown(clean))
                p.paragraph_format.left_indent = Inches(0.2)
                continue

            doc.add_paragraph(strip_markdown(stripped))

        buf = io.BytesIO()
        doc.save(buf)
        return buf.getvalue()
    except ImportError:
        return text.encode("utf-8")


def text_to_pdf_bytes(text: str, template_name: str = "Classic Professional") -> bytes:
    """Generate a styled PDF matching the selected template."""
    try:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle
        from reportlab.lib.units import cm
        from reportlab.lib import colors
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, HRFlowable
        from reportlab.lib.enums import TA_CENTER, TA_LEFT

        # ── Palette per template ──────────────────────────────────────
        PALETTES = {
            "Classic Professional": {"header": "#0d1b4b", "accent": "#0d1b4b",  "body": "#2d3748", "meta": "#4a5568"},
            "Modern Minimal":       {"header": "#1e3a8a", "accent": "#2563eb",  "body": "#374151", "meta": "#6b7280"},
            "Executive Elite":      {"header": "#1e293b", "accent": "#c9a96e",  "body": "#2d3748", "meta": "#6b5f4e"},
            "Tech Focused":         {"header": "#0f172a", "accent": "#059669",  "body": "#1e293b", "meta": "#64748b"},
            "Creative Clean":       {"header": "#4c1d95", "accent": "#7c3aed",  "body": "#374151", "meta": "#6b7280"},
            "Sleek Harvard":        {"header": "#000000", "accent": "#000000",  "body": "#222222", "meta": "#444444"},
            "Silicon Valley":       {"header": "#0f172a", "accent": "#0284c7",  "body": "#334155", "meta": "#64748b"},
            "Corporate Executive":  {"header": "#1e293b", "accent": "#1e293b",  "body": "#334155", "meta": "#475569"},
        }
        p = PALETTES.get(template_name, PALETTES["Classic Professional"])
        HDR  = colors.HexColor(p["header"])
        ACC  = colors.HexColor(p["accent"])
        BODY = colors.HexColor(p["body"])
        META = colors.HexColor(p["meta"])

        buf = io.BytesIO()
        doc = SimpleDocTemplate(
            buf, pagesize=A4,
            leftMargin=1.5*cm, rightMargin=1.5*cm,
            topMargin=1.2*cm, bottomMargin=1.2*cm,
        )

        name_style = ParagraphStyle("Name", fontName="Helvetica-Bold", fontSize=18,
            textColor=HDR, alignment=TA_CENTER, spaceAfter=2, leading=22)
        contact_style = ParagraphStyle("Contact", fontName="Helvetica", fontSize=8.5,
            textColor=META, alignment=TA_CENTER, spaceAfter=4, leading=11)
        section_style = ParagraphStyle("Section", fontName="Helvetica-Bold", fontSize=9,
            textColor=ACC, spaceBefore=8, spaceAfter=2, leading=12)
        job_header_style = ParagraphStyle("JobHdr", fontName="Helvetica-Bold", fontSize=10,
            textColor=HDR, spaceBefore=4, spaceAfter=1, leading=12)
        job_meta_style = ParagraphStyle("JobMeta", fontName="Helvetica", fontSize=8.5,
            textColor=META, spaceAfter=2, leading=11)
        body_style = ParagraphStyle("Body", fontName="Helvetica", fontSize=9,
            textColor=BODY, leading=12, spaceAfter=2)
        bullet_style = ParagraphStyle("Bullet", fontName="Helvetica", fontSize=9,
            textColor=BODY, leading=12, spaceAfter=1, leftIndent=14, firstLineIndent=-10)

        story = []
        lines = text.split("\n")
        name_done = False
        in_header = True

        for line in lines:
            stripped = line.strip()
            if not stripped:
                story.append(Spacer(1, 2))
                continue

            # Check for section header
            is_hdr, hdr_title = clean_and_check_header(stripped)
            if is_hdr:
                in_header = False
                story.append(Spacer(1, 3))
                story.append(HRFlowable(width="100%", thickness=0.8, color=ACC, spaceAfter=2))
                story.append(Paragraph(md_to_reportlab(hdr_title), section_style))
                continue

            # Name
            if in_header and not name_done:
                name_candidate = strip_markdown(stripped).strip()
                if len(name_candidate.split()) <= 6 and not any(c in name_candidate for c in ["@", "|", "+"]):
                    story.append(Paragraph(md_to_reportlab(name_candidate), name_style))
                    name_done = True
                    continue

            # Contact header lines
            if in_header:
                story.append(Paragraph(md_to_reportlab(stripped), contact_style))
                continue

            # Bullet
            if is_bullet(stripped):
                clean = clean_bullet_text(stripped)
                story.append(Paragraph(f"• {md_to_reportlab(clean)}", bullet_style))
                continue

            # Job header with |
            if "|" in stripped and len(stripped.split("|")) >= 2:
                parts = [p2.strip() for p2 in stripped.split("|")]
                company = md_to_reportlab(strip_markdown(parts[0]))
                story.append(Paragraph(company, job_header_style))
                if len(parts) > 1:
                    meta = " · ".join(md_to_reportlab(strip_markdown(p)) for p in parts[1:])
                    story.append(Paragraph(meta, job_meta_style))
                continue

            # Date line
            if re.search(r'\d{4}', stripped) and ("–" in stripped or "-" in stripped or "Present" in stripped) and len(stripped) < 60:
                story.append(Paragraph(md_to_reportlab(strip_markdown(stripped)), job_meta_style))
                continue

            story.append(Paragraph(md_to_reportlab(stripped), body_style))

        doc.build(story)
        return buf.getvalue()

    except ImportError:
        return text.encode("utf-8")
